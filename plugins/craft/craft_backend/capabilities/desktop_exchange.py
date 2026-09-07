"""Bounded typed spreadsheet actions over owned ArtifactRefs and Feishu identity."""
import base64
import csv
import io
import itertools
import json
import zipfile
from jsonschema import Draft202012Validator
from backend.capability_v2.atomic_web_contracts import obj
from backend.capability_v2.contracts import ArtifactRef, ExposurePolicy, BusinessInvariantContract
from backend.capability_v2.descriptor_adapter import descriptor_from_provider_spec
from backend.capability_v2.provider_contracts import CapabilitySpec, CapabilityBusinessError
from backend.platform_sdk.artifacts import read_artifact, create_artifact
from backend.platform_sdk.feishu import user_credential
from .data_exchange import _export_excel, _export_diff_report, _export_diff_lark_sheet
from .lark_exchange import read_lark_data, write_lark_data
from .desktop_pictures import resolve_picture

TEXT={'type':'string','maxLength':4096}
ID={'type':'string','minLength':1,'maxLength':128,'pattern':'^[A-Za-z0-9_.-]+$'}
SCALAR={'type':['string','number','boolean','null'],'maxLength':4096}
def array(schema,maximum,minimum=0):return {'type':'array','items':schema,'maxItems':maximum,'minItems':minimum}
COLUMN=obj({'key':TEXT,'label':TEXT,'width':{'type':'number','minimum':1,'maximum':1000},'include':{'type':'boolean'}},('key',))
COLUMNS=array(COLUMN,200,1)
ROWS=array(array(SCALAR,200),5000)
STYLE=obj({**{key:TEXT for key in ('headerBg','headerFg','altRowBg','borderStyle')},'fontSize':{'type':'integer','minimum':6,'maximum':48}})
DIFF_ROW=obj({'status':{'enum':['added','removed','modified','same']},'values_a':array(SCALAR,200),'values_b':array(SCALAR,200),'changed_fields':array(TEXT,200)},('status','values_a','values_b','changed_fields'))
DIFF=obj({'columns':COLUMNS,'diff_rows':array(DIFF_ROW,5000),'label_a':TEXT,'label_b':TEXT,'filename':TEXT},('columns','diff_rows'))
ARTIFACT=obj({'artifact_ref':ArtifactRef.model_json_schema(),'name':{'type':'string','minLength':1,'maxLength':255}},('artifact_ref','name'))
PARSED=obj({'headers':array(TEXT,200),'rows':array(array(TEXT,200),5000),'warnings':array(TEXT,20),'import_preview':obj({'import_preview_gid':ID,'content_hash':TEXT,'entry_count':{'type':'integer','minimum':0},'expires_at':TEXT})},('headers','rows','warnings'))

def fail(message): raise CapabilityBusinessError('invalid_input',message)

def parse(payload,context):
    data=read_artifact(payload['artifact_ref'],context,maximum=5*1024*1024)
    name=payload['name'].lower()
    workbook=None
    if name.endswith('.csv'):
        for encoding in ('utf-8-sig','gbk','latin-1'):
            try: text=data.decode(encoding);break
            except UnicodeDecodeError: pass
        source=csv.reader(io.StringIO(text))
    elif name.endswith(('.xlsx','.xlsm')):
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist())>50*1024*1024: fail('The workbook expands beyond 50 MiB.')
        workbook=openpyxl.load_workbook(io.BytesIO(data),read_only=True,keep_vba=False,data_only=True)
        sheet=workbook.active
        if sheet.max_column>200 or sheet.max_row>5002: fail('The workbook exceeds 5000 rows or 200 columns.')
        source=sheet.iter_rows(values_only=True)
    else: fail('Select a CSV, XLSX or XLSM document.')
    try:
        all_rows=list(itertools.islice(source,5002))
        if len(all_rows)>5001 or any(len(row)>200 for row in all_rows): fail('The document exceeds 5000 rows or 200 columns.')
        all_rows=[row for row in all_rows if any(value is not None and str(value).strip() for value in row)]
        if not all_rows: return {'success':True,'data':{'headers':[],'rows':[],'warnings':['文件为空']}}
        header=all_rows[0]
        if not name.endswith('.csv'):
            while header and (header[-1] is None or not str(header[-1]).strip()): header=header[:-1]
        headers=[str(value).strip() if value is not None and str(value).strip() else f'列{index+1}' for index,value in enumerate(header)]
        rows=[[str(value).strip() if value is not None else '' for value in [*row[:len(headers)],*['']*max(0,len(headers)-len(row))]] for row in all_rows[1:]]
        result={'headers':headers,'rows':rows,'warnings':[] if rows else ['未找到有效数据行（表头之外无数据）']}
        if payload.get('module')=='bop':
            from .bop_writes import import_preview
            result['import_preview']=import_preview({'document':{'version_tag':payload['name'].rsplit('.',1)[0],'bop_name':payload['name'],'entries':[dict(zip(headers,row)) for row in rows]}},context).data
        return {'success':True,'data':result}
    finally:
        if workbook is not None: workbook.close()

def diff_payload(payload):
    keys=[column['key'] for column in payload['columns']]
    return {**payload,'diff_rows':[{'_diffStatus':row['status'],'_rowA':dict(zip(keys,row['values_a'])),'_rowB':dict(zip(keys,row['values_b'])),'_changedFields':row['changed_fields']} for row in payload['diff_rows']]}

def exported(raw,context):
    return {'artifact_ref':create_artifact(base64.b64decode(raw['file_b64'],validate=True),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',context),'name':raw['filename']}

def export_excel(payload,context):
    columns=payload['columns'];keys=[column['key'] for column in columns]
    return exported(_export_excel({'template_config':{'columns':columns,'styles':payload.get('styles',{})},'rows':[dict(zip(keys,row)) for row in payload['rows']],'filename':payload.get('filename','')},context),context)

def export_diff(payload,context): return exported(_export_diff_report(diff_payload(payload),context),context)

def write_diff(payload,context):
    return _export_diff_lark_sheet({**diff_payload(payload),'user_access_token':user_credential(context)},context)

def write_sheet(payload,context):
    return write_lark_data({'operation':'sheets.write',**payload,'values':[payload['headers'],*payload['rows']],'user_access_token':user_credential(context)},context)['data']

def read_sheet(payload,context):
    return read_lark_data({'operation':'sheets.read',**payload,'user_access_token':user_credential(context)},context)['data']

def read_bitable(payload,context):
    return read_lark_data({'operation':'bitable.read',**payload,'user_access_token':user_credential(context)},context)['data']

def write_bitable(payload,context):
    value=write_lark_data({'operation':'bitable.write','app_token':payload['app_token'],'table_id':payload['table_id'],'records':[dict(zip(payload['headers'],row)) for row in payload['rows']],'user_access_token':user_credential(context)},context)['data']
    return {'success':value['success'],'written_rows':value['written_records']}

DEFINITIONS=[
 ('craft.bop.picture.resolve',obj({'version_gid':ID,'reference_hash':{'type':'string','pattern':'^[a-f0-9]{64}$'}},('version_gid','reference_hash')),ARTIFACT,False,resolve_picture),
 ('craft.data_exchange.feishu_bitable.read',obj({'app_token':ID,'table_id':ID,'page_size':{'type':'integer','minimum':1,'maximum':500}},('app_token','table_id')),PARSED,False,read_bitable),
 ('craft.data_exchange.feishu_bitable.write',obj({'app_token':ID,'table_id':ID,'headers':array(TEXT,200,1),'rows':ROWS},('app_token','table_id','headers','rows')),obj({'success':{'const':True},'written_rows':{'type':'integer','minimum':0,'maximum':5000}},('success','written_rows')),True,write_bitable),
 ('craft.data_exchange.excel.parse',obj({'artifact_ref':ArtifactRef.model_json_schema(),'name':ARTIFACT['properties']['name'],'module':TEXT},('artifact_ref','name')),obj({'success':{'const':True},'data':PARSED},('success','data')),False,parse),
 ('craft.data_exchange.excel.export',obj({'columns':COLUMNS,'rows':ROWS,'styles':STYLE,'filename':TEXT},('columns','rows')),ARTIFACT,True,export_excel),
 ('craft.data_exchange.diff_report.export',DIFF,ARTIFACT,True,export_diff),
 ('craft.data_exchange.feishu_diff.write',obj({**{key:value for key,value in DIFF['properties'].items() if key!='filename'},'spreadsheet_token':ID,'sheet_id':ID},('columns','diff_rows','spreadsheet_token','sheet_id')),obj({'written_rows':{'type':'integer','minimum':0,'maximum':5000},'spreadsheet_token':ID,'sheet_id':ID},('written_rows','spreadsheet_token','sheet_id')),True,write_diff),
 ('craft.data_exchange.feishu_sheet.write',obj({'spreadsheet_token':ID,'sheet_id':ID,'headers':array(TEXT,200,1),'rows':ROWS},('spreadsheet_token','sheet_id','headers','rows')),obj({'success':{'const':True},'written_rows':{'type':'integer','minimum':0,'maximum':5000}},('success','written_rows')),True,write_sheet),
 ('craft.data_exchange.feishu_sheet.read',obj({'spreadsheet_token':ID,'sheet_range':{**TEXT,'pattern':'^[A-Za-z0-9_ -]+![A-Z]+[0-9]+:[A-Z]+[0-9]+$'}},('spreadsheet_token','sheet_range')),PARSED,False,read_sheet),
]

def register_desktop_exchange(registry):
    for capability_id,input_schema,output_schema,write,service in DEFINITIONS:
        def handler(payload,context,_input=input_schema,_output=output_schema,_service=service):
            if not Draft202012Validator(_input).is_valid(payload): fail('The request does not match the spreadsheet action.')
            result=_service(payload,context)
            if not Draft202012Validator(_output).is_valid(result): raise CapabilityBusinessError('provider_error','The spreadsheet outcome exceeds its closed model.')
            return result
        effect=('Write ' if write else 'Read ')+capability_id.removeprefix('craft.data_exchange.').replace('.',' ')+' using owned immutable artifacts or the authenticated Feishu delegation.'
        spec=CapabilitySpec(id=capability_id,version=1,owner='craft',description=effect,use_when=effect,do_not_use_when='Another export, import, resource or remote method is requested.',risk='write' if write else 'read',confirmation='user' if write else 'none',idempotent=True,permissions=('craft.write',) if write else ('craft.read',),input_schema=input_schema,output_schema=output_schema,tags=('craft','desktop','artifact'))
        descriptor=descriptor_from_provider_spec(spec).model_copy(update={'business_effect':effect,'business_acceptance_criteria':(effect,'At most 5000 rows and 200 columns are accepted; raw credentials and filesystem paths are never input.','File output is an immutable ArtifactRef with verified hash and size.'),'business_invariants':(BusinessInvariantContract(rule_id=capability_id+'.owner_bound',version=1,statement='File and Feishu access are derived from authenticated context.',applies_when='The data exchange action executes.',enforcement_ref='backend/platform_sdk/artifacts.py; backend/platform_sdk/feishu.py',error_code='permission_denied',test_refs=('backend/tests/test_desktop_round5_exchange.py',)),),'no_business_invariant_reason':None,'exposure':ExposurePolicy(web=True,api=True,plugin=False,agent=False,mcp=False),'consistency_policy':'external' if write else 'strong','transaction_policy':{'mode':'provider','boundary':'owning_domain'},'delegation_policy':'none','idempotency_policy':'required' if write else 'none','evidence_policy':'optional'})
        registry.register(spec,handler,descriptor=descriptor)
